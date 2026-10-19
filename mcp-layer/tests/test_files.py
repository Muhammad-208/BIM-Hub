import json

import pytest

from bimhub.cli.main import main
from bimhub.core.errors import BimHubError
from bimhub.core.paths import HubPaths
from bimhub.core.project import Project
from bimhub.files import sources as S
from bimhub.files.readers import read_file

IFC = """ISO-10303-21;
HEADER;
FILE_DESCRIPTION(('ViewDefinition [CoordinationView]'),'2;1');
FILE_NAME('TST-WYA-ZZ-ZZ-M3-S-0001.ifc','2026-10-16',(''),(''),'','Autodesk Revit 2027','');
FILE_SCHEMA(('IFC4'));
ENDSEC;
DATA;
#1=IFCBUILDINGSTOREY('2xyz',#5,'Level 03',$,$,#9,$,$,.ELEMENT.,9600.);
#2=IFCBEAM('3abc',#5,'B3-12',$,$,$,$,$,$);
#3=IFCBEAM('3abd',#5,'B3-15',$,$,$,$,$,$);
#4=IFCDUCTSEGMENT('4abc',#5,'Duct',$,$,$,$,$,$);
ENDSEC;
END-ISO-10303-21;
"""


@pytest.fixture()
def acc(hub, tmp_path):
    root = tmp_path / "ACCDocs" / "WYA" / "TST" / "Project Files"
    (root / "01 WIP" / "STR").mkdir(parents=True)
    (root / "01 WIP" / "STR" / "TST-WYA-ZZ-ZZ-M3-S-0001.ifc").write_text(IFC)
    (root / "01 WIP" / "STR" / "Backup").mkdir()
    (root / "01 WIP" / "STR" / "Backup" / "old.ifc").write_text(IFC)
    (root / "00 Client").mkdir()
    (root / "00 Client" / "EIR notes.txt").write_text("Exchange information requirements: COBie at handover.")
    (root / "00 Client" / "rooms.csv").write_text("Room,Area\n101,24\n102,30\n")
    (root / "TST-WYA-ZZ-ZZ-M3-M-0001.nwc").write_bytes(b"\x00NWC" + b"\x00" * 300)
    (root / "site.dwg").write_bytes(b"AC1032" + b"\x00" * 50)
    import docx, openpyxl
    d = docx.Document(); d.add_heading("Employer Information Requirements", 1); d.add_paragraph("LOIN per stage.")
    t = d.add_table(rows=1, cols=2); t.cell(0, 0).text = "Stage"; t.cell(0, 1).text = "LOD"
    d.save(root / "00 Client" / "EIR.docx")
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "MIDP"; ws.append(["Container", "Due"]); ws.append(["TST-WYA-ZZ-03-DR-S-0101", "2026-11-01"])
    wb.save(root / "00 Client" / "MIDP.xlsx")
    assert main(["source", "add", "TST", "--name", "ACC", "--type", "acc-dc", "--path", str(root)]) == 0
    assert main(["source", "add", "TST", "--name", "ACC-web", "--type", "acc-web", "--url", "https://acc.autodesk.com/x"]) == 0
    return root


def test_sources_scan_find_read(hub, acc):
    p = Project.load(HubPaths.discover())
    names = [s.name for s in S.sources(p)]
    assert names == ["project", "ACC", "ACC-web"]
    text = (p.dir / "project.yaml").read_text(encoding="utf-8")
    assert "# Where this project's files live" in text           # comments survive the rewrite
    r = S.scan(p)
    kinds = r["by_kind"]
    assert kinds["ifc"] == 1                                      # Backup/ excluded
    assert kinds["word"] == 1 and kinds["excel"] == 1 and kinds["navisworks-cache"] == 1 and kinds["dwg"] == 1
    assert any(n.get("source") == "ACC-web" for n in r["notes"])
    f = S.find(p, kind="ifc")
    assert f["files"][0]["discipline"] == "str" and f["files"][0]["iso_name_ok"]
    ifc = read_file(S.resolve(p, "ACC", f["files"][0]["path"]))
    assert ifc["schema"] == "IFC4" and ifc["storeys"] == ["Level 03"] and ifc["elements_by_type"]["IFCBEAM"] == 2
    assert "COBie" in read_file(S.resolve(p, "ACC", "00 Client/EIR notes.txt"))["text"]
    dx = read_file(S.resolve(p, "ACC", "00 Client/EIR.docx"))
    assert dx["text"].startswith("# Employer") and dx["tables"][0][0] == ["Stage", "LOD"]
    xl = read_file(S.resolve(p, "ACC", "00 Client/MIDP.xlsx"))
    assert xl["sheet"] == "MIDP" and xl["rows"][1][0].startswith("TST-")
    assert read_file(S.resolve(p, "ACC", "site.dwg"))["autocad_format"] == "2018"
    assert "navisworks" in read_file(S.resolve(p, "ACC", "TST-WYA-ZZ-ZZ-M3-M-0001.nwc"))["open_with"]
    with pytest.raises(BimHubError):
        S.resolve(p, "ACC", "../../../etc/passwd")
    with pytest.raises(BimHubError):
        S.resolve(p, "ACC-web", "anything")


def test_hub_tools(hub, acc):
    from bimhub.servers import hub as srv
    srv.ctx._project = None
    srv.ctx.paths = HubPaths.discover()
    assert srv.files_scan()["count"] >= 6
    hits = srv.files_find(query="EIR", source="ACC")
    assert hits["matched"] == 2
    assert srv.file_read("ACC", "00 Client/rooms.csv")["total_rows"] == 2
    assert srv.file_read("ACC", "../x")["ok"] is False
