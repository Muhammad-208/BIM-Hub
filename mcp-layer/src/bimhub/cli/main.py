"""bimhub CLI — manage projects, teams, the Learning Hub library and MCP wiring.

  bimhub new-project <CODE> --name "..." --team arc,str,mech
  bimhub sync [--project CODE]          regenerate agents/skills/.mcp.json for one or all projects
  bimhub list                           projects and their teams
  bimhub catalog [--root PATH ...]      rebuild learning-hub/library/catalog.yaml from your source folders
  bimhub ingest <ID> [<ID> ...]         extract a catalogued file to learning-hub/knowledge/_raw/<ID>.md
  bimhub doctor                         check config, projects and desktop bridges
  bimhub source add AOH --name ACC --type acc-dc --path "C:/Users/me/DC/ACCDocs/WYA/AOH/Project Files"
  bimhub scan AOH                       index every file in the project's sources
"""
from __future__ import annotations

import argparse
import fnmatch
import hashlib
import json
import re
import shutil
import sys
from collections import defaultdict
from pathlib import Path

import yaml

from ..core.config import load_hub_settings, load_servers
from ..core.errors import BimHubError
from ..core.paths import HubPaths
from ..core.project import CORE_AGENTS, ROLE_TITLES, ROLES, Project, split_agent
from ..core import project_types as ptypes

TEMPLATED = {".md", ".yaml", ".yml"}


def _render(text: str, code: str, name: str, team: list[str], ptype: dict | None = None) -> str:
    text = (text.replace("{{CODE}}", code).replace("{{NAME}}", name)
                .replace("{{TEAM}}", ", ".join(team)))
    if "{{TYPE" in text or "{{ROLES}}" in text or "{{MCP}}" in text or "{{PARAMS}}" in text:
        t = ptype or {}
        params = yaml.safe_dump(t.get("params") or {}, allow_unicode=True, sort_keys=False, default_flow_style=False)
        params = "\n".join("  " + ln for ln in params.strip().splitlines()) if t.get("params") else "  {}"
        text = (text.replace("{{TYPE}}", t.get("id", "")).replace("{{TYPE_TITLE}}", t.get("title", "custom"))
                    .replace("{{ROLES}}", ", ".join(t.get("roles") or ROLES))
                    .replace("{{MCP}}", ", ".join(t.get("mcp") or ["bimhub", "revit", "navisworks"]))
                    .replace("{{PARAMS}}", params.rstrip() if t.get("params") else "  {}"))
    return text


def _agent_templates(paths: HubPaths) -> dict[str, Path]:
    """Core agents (L2–L4): one .md each. Discipline roles live in L1-disciplines/roles/."""
    return {p.stem: p for p in paths.agents.rglob("*.md")
            if "_legacy" not in p.parts and "L1-disciplines" not in p.parts}


def _role_templates(paths: HubPaths) -> dict[str, str]:
    d = paths.agents / "L1-disciplines" / "roles"
    return {r: (d / f"{r}.md").read_text(encoding="utf-8") for r in ROLES if (d / f"{r}.md").is_file()}


def _discipline_profiles(paths: HubPaths) -> dict[str, dict]:
    f = paths.agents / "L1-disciplines" / "disciplines.yaml"
    return (yaml.safe_load(f.read_text(encoding="utf-8")) or {}).get("disciplines", {}) if f.is_file() else {}


def _profile_md(disc: str, prof: dict) -> str:
    return (f"- **Scope:** {prof.get('scope', '')}.\n"
            f"- **Revit / authoring categories:** {prof.get('categories', '')}.\n"
            f"- **Information (LOI) owned:** {prof.get('loi', '')}. Exact properties per stage: LOIN table in `01-bep/`.\n"
            f"- **Coordination hotspots:** {prof.get('hotspots', '')}.\n"
            f"- **Codes & references:** {prof.get('codes', '')}. Hub: `knowledge/02-disciplines/{disc}/` first; cite what you use.\n"
            f"- Project-specific notes: `team-notes/{disc}.md` (read it before starting).")


def _render_role(text: str, p: Project, disc: str, prof: dict) -> str:
    sub_team = ", ".join(f"`{disc}-{r}` ({ROLE_TITLES[r]})" for r in p.roles)
    text = (text.replace("{{DISC_UPPER}}", disc.upper()).replace("{{DISC_TITLE}}", prof.get("title", disc.upper()))
                .replace("{{DISC}}", disc).replace("{{ROLE_CODE}}", str((p.disciplines.get(disc) or {}).get("role")
                                                                         or prof.get("role_code", "?")))
                .replace("{{SUBTEAM}}", sub_team).replace("{{PROFILE}}", _profile_md(disc, prof)))
    return _render(text, p.code, p.name, p.team)


# ───────────────────────── new-project ─────────────────────────
def cmd_new_project(paths: HubPaths, a: argparse.Namespace) -> None:
    code = a.code.upper()
    if not re.fullmatch(r"[A-Z0-9]{2,6}", code):
        raise BimHubError("Project code must be 2–6 letters/digits (it is the first field of every file name)")
    dst = paths.projects / code
    if dst.exists():
        raise BimHubError(f"{dst} already exists")
    profiles = _discipline_profiles(paths)
    team = [t.strip().lower() for t in a.team.split(",") if t.strip()]
    unknown = [t for t in team if t not in profiles]
    if unknown:
        raise BimHubError(f"Unknown discipline(s) {unknown}. Available: {', '.join(profiles)}")
    ptype = ptypes.get_type(paths, getattr(a, "type", None)) if getattr(a, "type", None) else None
    shutil.copytree(paths.template, dst)
    for f in dst.rglob("*"):
        if f.is_file() and f.suffix in TEMPLATED:
            f.write_text(_render(f.read_text(encoding="utf-8"), code, a.name, team, ptype), encoding="utf-8")
    print(f"Created project {code} at {dst}" + (f" — type {ptype['id']} ({ptype.get('title')})" if ptype else ""))
    sync_project(paths, code)


# ───────────────────────── sync ─────────────────────────
def _mcp_json(paths: HubPaths, p: Project) -> dict:
    servers = load_servers(paths)
    out = {}
    for name in p.mcp:
        s = servers.get(name)
        if not s:
            print(f"  ! {p.code}: MCP '{name}' is not in servers.yaml — skipped", file=sys.stderr)
            continue
        env = {"BIMHUB_ROOT": str(paths.root), "BIMHUB_PROJECT": p.code, **s.env}
        if s.kind == "native":
            out[name] = {"command": sys.executable, "args": ["-m", s.module], "env": env}
        else:
            out[name] = {"command": s.command, "args": s.args, "env": {**s.env}}
    return {"mcpServers": out}


def sync_project(paths: HubPaths, code: str) -> None:
    p = Project.load(paths, code)
    templates = _agent_templates(paths)
    agents_dir = p.dir / ".claude" / "agents"
    skills_dir = p.dir / ".claude" / "skills"
    if agents_dir.exists():
        shutil.rmtree(agents_dir)          # generated — project-specific notes live in team-notes/
    agents_dir.mkdir(parents=True)
    roles, profiles = _role_templates(paths), _discipline_profiles(paths)
    for agent in p.all_agents:
        disc, role = split_agent(agent)
        if disc:
            if role not in roles or disc not in profiles:
                print(f"  ! {code}: no role template / profile for '{agent}'", file=sys.stderr)
                continue
            text = _render_role(roles[role], p, disc, profiles[disc])
        else:
            src = templates.get(agent)
            if not src:
                print(f"  ! {code}: no template for agent '{agent}'", file=sys.stderr)
                continue
            text = _render(src.read_text(encoding="utf-8"), p.code, p.name, p.team)
        (agents_dir / f"{agent}.md").write_text(text, encoding="utf-8")
    if skills_dir.exists():
        shutil.rmtree(skills_dir)
    shutil.copytree(paths.skills, skills_dir)
    (p.dir / ".mcp.json").write_text(json.dumps(_mcp_json(paths, p), indent=2), encoding="utf-8")
    settings = p.dir / ".claude" / "settings.json"
    if not settings.exists():
        settings.write_text(json.dumps({"enableAllProjectMcpServers": True}, indent=2), encoding="utf-8")
    for disc in p.team:
        wip = (p.disciplines.get(disc) or {}).get("wip") or f"02-cde/WIP/{disc.upper()}"
        (p.dir / wip).mkdir(parents=True, exist_ok=True)
    t = ptypes.get_type(paths, p.type_id) if p.type_id else None
    wf = p.dir / "WORKFLOW.md"
    if t:
        for d in t.get("folders") or []:
            (p.dir / d).mkdir(parents=True, exist_ok=True)
        wf.write_text(ptypes.render_workflow_md(p, t), encoding="utf-8")
        cm = p.dir / "CLAUDE.md"
        txt = cm.read_text(encoding="utf-8") if cm.is_file() else f"# Project {p.code}\n"
        if "@WORKFLOW.md" not in txt:
            cm.write_text(txt.rstrip() + "\n\n## Workflow for this project type\n@WORKFLOW.md\n", encoding="utf-8")
        missing = [s for s in t.get("skills") or [] if not (paths.skills / s).is_dir()]
        if missing:
            print(f"  ! {code}: type {t['id']} needs skills not in the hub: {missing}", file=sys.stderr)
    elif not wf.exists():
        wf.write_text("# Workflow\n\nNo project type set — free workflow. Set `type:` in project.yaml "
                      "(see `bimhub types`) and run `bimhub sync`.\n", encoding="utf-8")
    print(f"Synced {code}: {len(p.all_agents)} agents, {len(list(paths.skills.iterdir()))} skills, "
          f"MCP: {', '.join(p.mcp)}" + (f", type: {t['id']}" if t else ", type: (none)"))


def cmd_types(paths: HubPaths, a: argparse.Namespace) -> None:
    for tid, t in ptypes.load_types(paths).items():
        print(f"{tid:14} {t.get('title')}  —  roles: {', '.join(t.get('roles') or ROLES)}; "
              f"stages: {' → '.join(s['id'] for s in t.get('stages') or [])}")


def cmd_sync(paths: HubPaths, a: argparse.Namespace) -> None:
    codes = [a.project.upper()] if a.project else paths.list_projects()
    if not codes:
        print("No projects yet. Create one with: bimhub new-project <CODE> --name ... --team ...")
    for c in codes:
        sync_project(paths, c)


def cmd_list(paths: HubPaths, a: argparse.Namespace) -> None:
    for c in paths.list_projects():
        p = Project.load(paths, c)
        print(f"{c:8} {p.name:30} type: {p.type_id or '-':13} team: {', '.join(p.team)}")


# ───────────────────────── catalog / ingest ─────────────────────────
YEAR = re.compile(r"(?<!\d)(19[89]\d|20[0-3]\d)(?!\d)")


def cmd_catalog(paths: HubPaths, a: argparse.Namespace) -> None:
    settings = load_hub_settings(paths)
    roots = [Path(r) for r in (a.root or settings.get("library_roots", []))]
    ignore = settings.get("ignore_patterns", [])
    rules = settings.get("tag_rules", {})
    items, by_name, by_size = [], defaultdict(list), defaultdict(list)
    for root in roots:
        if not root.is_dir():
            print(f"  ! library root not found: {root}", file=sys.stderr)
            continue
        for f in sorted(root.rglob("*")):
            if not f.is_file() or any(fnmatch.fnmatch(f.name, pat) for pat in ignore):
                continue
            rel = f.relative_to(root).as_posix()
            low = rel.lower()
            tags = sorted(t for t, kws in rules.items() if any(k in low for k in kws))
            years = YEAR.findall(f.stem)
            item = {"id": hashlib.sha1(f"{root.name}/{rel}".encode()).hexdigest()[:8],
                    "title": re.sub(r"[_]+", " ", f.stem)[:140], "root": root.as_posix(), "path": rel,
                    "type": f.suffix.lower().lstrip("."), "size_kb": f.stat().st_size // 1024,
                    "year": int(max(years)) if years else None, "tags": tags, "status": "current"}
            items.append(item)
            by_name[f.name.lower()].append(item)
            by_size[(f.suffix.lower(), f.stat().st_size)].append(item)
    for group in list(by_name.values()) + list(by_size.values()):
        for dup in group[1:]:
            if dup.get("duplicate_of") is None and dup["id"] != group[0]["id"]:
                dup["duplicate_of"] = group[0]["id"]
    # LOD spec family: newest year is current, older are superseded
    lod = [i for i in items if "lod-loin" in i["tags"] and "spec" in i["title"].lower() and i["year"]]
    if lod:
        newest = max(i["year"] for i in lod)
        for i in lod:
            if i["year"] < newest:
                i["status"] = "superseded"
    out = paths.hub / "library" / "catalog.yaml"
    out.parent.mkdir(parents=True, exist_ok=True)
    summary = defaultdict(int)
    for i in items:
        for t in i["tags"] or ["untagged"]:
            summary[t] += 1
    with open(out, "w", encoding="utf-8") as fh:
        fh.write("# Generated by `bimhub catalog` — re-run after adding files. Edit tags/status here if needed.\n")
        yaml.safe_dump({"count": len(items), "by_tag": dict(sorted(summary.items())), "items": items},
                       fh, allow_unicode=True, sort_keys=False, width=200)
    dups = sum(1 for i in items if i.get("duplicate_of"))
    print(f"Catalogued {len(items)} files ({dups} duplicates flagged) → {out}")


def _extract(path: Path) -> str:
    ext = path.suffix.lower()
    if ext in (".md", ".txt"):
        return path.read_text(encoding="utf-8", errors="replace")
    try:
        if ext == ".pdf":
            try:
                import pymupdf as fitz
            except ImportError:
                import fitz  # older pymupdf
            with fitz.open(path) as doc:
                return "\n\n".join(f"<!-- page {i+1} -->\n{pg.get_text()}" for i, pg in enumerate(doc))
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            lines = []
            for para in d.paragraphs:
                s = para.style.name.lower() if para.style is not None else ""
                prefix = "#" * int(s[-1]) + " " if s.startswith("heading") and s[-1:].isdigit() else ""
                lines.append(prefix + para.text)
            for t in d.tables:
                for r in t.rows:
                    lines.append("| " + " | ".join(c.text.strip() for c in r.cells) + " |")
            return "\n".join(lines)
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"## {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    if any(v is not None for v in row):
                        out.append("| " + " | ".join("" if v is None else str(v) for v in row) + " |")
            return "\n".join(out)
    except ImportError as e:
        raise BimHubError(f"Missing extractor ({e.name}). Install: pip install -e .[ingest]") from e
    raise BimHubError(f"No extractor for {ext} (videos need transcription first)")


def cmd_ingest(paths: HubPaths, a: argparse.Namespace) -> None:
    cat = yaml.safe_load((paths.hub / "library" / "catalog.yaml").read_text(encoding="utf-8"))
    index = {i["id"]: i for i in cat.get("items", [])}
    raw = paths.hub / "knowledge" / "_raw"
    raw.mkdir(parents=True, exist_ok=True)
    for cid in a.ids:
        item = index.get(cid)
        if not item:
            print(f"  ! {cid}: not in catalog", file=sys.stderr)
            continue
        src = Path(item["root"]) / item["path"]
        text = _extract(src)
        header = (f"---\ntitle: \"{item['title']}\"\ncatalog_id: {cid}\nsource: \"{src.as_posix()}\"\n"
                  f"status: raw\ntags: [{', '.join(item['tags'])}]\n---\n\n")
        (raw / f"{cid}.md").write_text(header + text, encoding="utf-8")
        print(f"Ingested {cid} → knowledge/_raw/{cid}.md ({len(text)//1000} k chars)")


# ───────────────────────── sources ─────────────────────────
def cmd_source(paths: HubPaths, a: argparse.Namespace) -> None:
    p = Project.load(paths, a.project.upper())
    f = p.dir / "project.yaml"
    data = yaml.safe_load(f.read_text(encoding="utf-8")) or {}
    srcs = data.get("sources") or []
    if a.action == "list":
        from ..files.sources import sources
        for s in sources(p):
            print(f"{s.name:14} {s.type:8} {'ok ' if s.readable else '-- '} {s.path or s.url}")
        return
    if a.action == "add":
        if not (a.path or a.url):
            raise BimHubError("--path or --url is required")
        if any(s.get("name", "").lower() == a.name.lower() for s in srcs):
            raise BimHubError(f"Source '{a.name}' already exists")
        entry = {"name": a.name, "type": a.type}
        if a.path:
            entry["path"] = a.path.replace("\\", "/")
            if not Path(a.path).is_dir():
                print(f"  ! warning: {a.path} is not a folder on this PC (yet)", file=sys.stderr)
        if a.url:
            entry["url"] = a.url
        srcs.append(entry)
    elif a.action == "remove":
        srcs = [s for s in srcs if s.get("name", "").lower() != a.name.lower()]
    _set_sources(f, srcs)
    print(f"{p.code}: {len(srcs)} source(s) — run: bimhub scan {p.code}")


def _set_sources(f: Path, srcs: list) -> None:
    """Rewrite only the `sources:` block so the comments elsewhere in project.yaml survive."""
    text = f.read_text(encoding="utf-8")
    block = yaml.safe_dump({"sources": srcs}, allow_unicode=True, sort_keys=False).rstrip() if srcs else "sources: []"
    lines, out, skip = text.splitlines(), [], False
    for ln in lines:
        if ln.startswith("sources:"):
            skip = True
            out.append(block)
            continue
        if skip and (ln.startswith((" ", "-", "\t")) or not ln.strip()):
            if not ln.strip():
                skip = False
                out.append(ln)
            continue
        skip = False
        out.append(ln)
    if not any(l.startswith("sources:") for l in lines):
        out += ["", block]
    f.write_text("\n".join(out) + "\n", encoding="utf-8")


def cmd_scan(paths: HubPaths, a: argparse.Namespace) -> None:
    from ..files.sources import scan
    r = scan(Project.load(paths, a.project.upper()), only=a.source)
    print(f"Indexed {r['count']} files → {r['index']}")
    for k, v in r["by_kind"].items():
        print(f"  {k:22} {v}")
    for n in r["notes"]:
        print(f"  ! {n}")


def cmd_studio(paths: HubPaths, a: argparse.Namespace) -> None:
    try:
        from ..studio.server import run
    except ImportError as e:
        raise BimHubError(f"Studio dependencies missing ({e.name}): pip install -e .[studio]") from e
    print(f"BIM Hub Studio → http://localhost:{a.port}  (Ctrl+C to stop)")
    run(a.port, not a.no_browser)


# ───────────────────────── doctor ─────────────────────────
def cmd_doctor(paths: HubPaths, a: argparse.Namespace) -> None:
    print(f"Hub root: {paths.root}")
    servers = load_servers(paths)
    print(f"MCP registry: {', '.join(f'{n} ({s.kind})' for n, s in servers.items())}")
    print(f"Agents: {len(_agent_templates(paths))} core + {len(_role_templates(paths))} roles × "
          f"{len(_discipline_profiles(paths))} disciplines · Skills: {len(list(paths.skills.iterdir()))}")
    print(f"Projects: {', '.join(paths.list_projects()) or '(none)'}")
    from ..core.bridge import make_bridge
    for n, s in servers.items():
        if s.bridge:
            try:
                h = make_bridge(s.bridge, n).health()
                print(f"  ✓ {n}: {h}")
            except Exception as e:  # noqa: BLE001
                print(f"  · {n}: not running ({s.bridge.url})")


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(prog="bimhub", description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("new-project"); s.add_argument("code"); s.add_argument("--name", required=True)
    s.add_argument("--team", required=True, help="comma list: arc,str,mech,plumb,ff,elec,ict,infra,land")
    s.add_argument("--type", help="project type: cad-to-bim | model-review | lod-uplift | scan-to-bim (bimhub types)")
    sub.add_parser("types", help="list project types and their workflows")
    s = sub.add_parser("sync"); s.add_argument("--project")
    sub.add_parser("list")
    s = sub.add_parser("catalog"); s.add_argument("--root", action="append")
    s = sub.add_parser("ingest"); s.add_argument("ids", nargs="+")
    sub.add_parser("doctor")
    s = sub.add_parser("source", help="list | add | remove project file sources (local, ACC Desktop Connector, ACC web)")
    s.add_argument("action", choices=["list", "add", "remove"]); s.add_argument("project")
    s.add_argument("--name"); s.add_argument("--type", default="local", choices=["local", "acc-dc", "acc-web"])
    s.add_argument("--path"); s.add_argument("--url")
    s = sub.add_parser("scan", help="index all files of a project's sources"); s.add_argument("project"); s.add_argument("--source")
    s = sub.add_parser("studio", help="start BIM Hub Studio (API + UI) on http://localhost:8800")
    s.add_argument("--port", type=int, default=8800); s.add_argument("--no-browser", action="store_true")
    a = ap.parse_args(argv)
    try:
        paths = HubPaths.discover()
        {"new-project": cmd_new_project, "types": cmd_types, "sync": cmd_sync, "list": cmd_list, "catalog": cmd_catalog,
         "ingest": cmd_ingest, "doctor": cmd_doctor, "source": cmd_source, "scan": cmd_scan, "studio": cmd_studio}[a.cmd](paths, a)
        return 0
    except BimHubError as e:
        print(f"error: {e}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
