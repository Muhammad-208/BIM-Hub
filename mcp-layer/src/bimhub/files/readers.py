"""Readers per file type. Each returns a dict the agent can use directly.

Text-like formats are read fully (truncated to max_chars). Binary BIM formats that only the authoring
app can open (RVT, NWC/NWD/NWF, DWG) return what can be read WITHOUT the app (version, worksharing,
central path…) plus `open_with`: the MCP server to use for full content.
"""
from __future__ import annotations

import csv
import io
import re
from collections import Counter
from pathlib import Path
from typing import Any, Callable

from ..core.errors import BimHubError

KIND = {
    ".rvt": "revit-model", ".rfa": "revit-family", ".rte": "revit-template",
    ".ifc": "ifc", ".ifczip": "ifc",
    ".nwc": "navisworks-cache", ".nwd": "navisworks-document", ".nwf": "navisworks-fileset",
    ".dwg": "dwg", ".dxf": "dxf",
    ".docx": "word", ".doc": "word-legacy", ".xlsx": "excel", ".xlsm": "excel", ".xls": "excel-legacy",
    ".csv": "csv", ".txt": "text", ".md": "text", ".json": "text", ".xml": "text", ".log": "text",
    ".pdf": "pdf", ".pptx": "powerpoint",
    ".png": "image", ".jpg": "image", ".jpeg": "image",
    ".bcf": "bcf", ".bcfzip": "bcf",
    ".e57": "point-cloud", ".rcp": "point-cloud", ".rcs": "point-cloud", ".las": "point-cloud",
    ".laz": "point-cloud", ".pts": "point-cloud", ".ptx": "point-cloud", ".xyz": "point-cloud",
}


def kind_of(path: Path) -> str:
    return KIND.get(path.suffix.lower(), "other")


def _cut(text: str, max_chars: int) -> dict[str, Any]:
    return {"text": text[:max_chars], "truncated": len(text) > max_chars, "chars": len(text)}


# ───────────── text-like ─────────────
def read_text(p: Path, max_chars: int, **_: Any) -> dict:
    return _cut(p.read_text(encoding="utf-8", errors="replace"), max_chars)


def read_csv(p: Path, max_chars: int, max_rows: int = 200, **_: Any) -> dict:
    with open(p, encoding="utf-8", errors="replace", newline="") as f:
        rows = list(csv.reader(f))
    return {"columns": rows[0] if rows else [], "rows": rows[1:max_rows + 1], "total_rows": max(0, len(rows) - 1)}


def read_docx(p: Path, max_chars: int, **_: Any) -> dict:
    try:
        import docx
    except ImportError as e:
        raise BimHubError("python-docx not installed: pip install -e .[ingest]") from e
    d = docx.Document(str(p))
    out = []
    for para in d.paragraphs:
        s = (para.style.name or "").lower() if para.style is not None else ""
        lvl = s[-1] if s.startswith("heading") and s[-1:].isdigit() else ""
        out.append(("#" * int(lvl) + " " if lvl else "") + para.text)
    tables = []
    for t in d.tables[:30]:
        tables.append([[c.text.strip() for c in r.cells] for r in t.rows[:60]])
    res = _cut("\n".join(out), max_chars)
    res["tables"] = tables
    return res


def read_xlsx(p: Path, max_chars: int, sheet: str | None = None, max_rows: int = 200, **_: Any) -> dict:
    try:
        import openpyxl
    except ImportError as e:
        raise BimHubError("openpyxl not installed: pip install -e .[ingest]") from e
    wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
    names = wb.sheetnames
    target = sheet or names[0]
    if target not in names:
        raise BimHubError(f"Sheet '{target}' not found. Sheets: {names}")
    ws = wb[target]
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i > max_rows:
            break
        rows.append(["" if v is None else str(v) for v in r])
    return {"sheets": names, "sheet": target, "rows": rows, "max_row": ws.max_row, "max_column": ws.max_column}


def read_pdf(p: Path, max_chars: int, pages: str | None = None, **_: Any) -> dict:
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz  # type: ignore
        except ImportError as e:
            raise BimHubError("pymupdf not installed: pip install -e .[ingest]") from e
    with fitz.open(p) as doc:
        n = doc.page_count
        sel = range(n)
        if pages:
            a, _, b = pages.partition("-")
            sel = range(max(0, int(a) - 1), min(n, int(b or a)))
        text = "\n\n".join(f"<!-- page {i + 1} -->\n{doc[i].get_text()}" for i in sel)
        res = _cut(text, max_chars)
        res["page_count"] = n
        return res


def read_pptx(p: Path, max_chars: int, **_: Any) -> dict:
    import zipfile
    texts = []
    with zipfile.ZipFile(p) as z:
        slides = sorted((n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)),
                        key=lambda n: int(re.findall(r"\d+", n)[-1]))
        for i, n in enumerate(slides, 1):
            xml = z.read(n).decode("utf-8", "replace")
            texts.append(f"## Slide {i}\n" + " ".join(re.findall(r"<a:t>([^<]*)</a:t>", xml)))
    return _cut("\n\n".join(texts), max_chars)


# ───────────── IFC (STEP text, streamed — no ifcopenshell needed) ─────────────
ENTITY = re.compile(r"^#\d+\s*=\s*(IFC[A-Z0-9_]+)\s*\(", re.I)
STOREY = re.compile(r"^#\d+\s*=\s*IFCBUILDINGSTOREY\('([^']*)',[^,]*,'([^']*)'", re.I)


def read_ifc(p: Path, max_chars: int, **_: Any) -> dict:
    if p.suffix.lower() == ".ifczip":
        import zipfile
        with zipfile.ZipFile(p) as z:
            inner = next((n for n in z.namelist() if n.lower().endswith(".ifc")), None)
            if not inner:
                raise BimHubError("No .ifc inside the ifczip")
            stream = io.TextIOWrapper(z.open(inner), encoding="utf-8", errors="replace")
            return _ifc_stream(stream)
    with open(p, encoding="utf-8", errors="replace") as f:
        return _ifc_stream(f)


def _ifc_stream(f: Any) -> dict:
    header, counts, storeys = [], Counter(), []
    in_header = True
    for line in f:
        if in_header:
            header.append(line.strip())
            if line.strip().upper().startswith("ENDSEC"):
                in_header = False
            continue
        m = ENTITY.match(line)
        if m:
            counts[m.group(1).upper()] += 1
            s = STOREY.match(line)
            if s:
                storeys.append(s.group(2))
    head = " ".join(header)
    schema = re.search(r"FILE_SCHEMA\s*\(\s*\(\s*'([^']+)'", head, re.I)
    app = re.search(r"FILE_NAME\s*\((.*?)\);", head, re.I)
    physical = {k: v for k, v in counts.items() if k in PHYSICAL}
    return {"schema": schema.group(1) if schema else None, "file_name_header": app.group(1)[:400] if app else None,
            "storeys": storeys, "elements_by_type": dict(sorted(physical.items(), key=lambda kv: -kv[1])),
            "entities_total": sum(counts.values()),
            "open_with": "revit (export_ifc to produce) · ifcopenshell for deep queries"}


PHYSICAL = {e.upper() for e in """IfcWall IfcWallStandardCase IfcSlab IfcBeam IfcColumn IfcDoor IfcWindow IfcStair IfcRamp IfcRoof
IfcCovering IfcCurtainWall IfcRailing IfcPlate IfcMember IfcFooting IfcPile IfcBuildingElementProxy IfcSpace IfcFurnishingElement
IfcDuctSegment IfcDuctFitting IfcPipeSegment IfcPipeFitting IfcCableCarrierSegment IfcCableCarrierFitting IfcFlowTerminal
IfcFlowSegment IfcFlowFitting IfcAirTerminal IfcSanitaryTerminal IfcLightFixture IfcFireSuppressionTerminal IfcPump IfcFan
IfcValve IfcElectricDistributionBoard IfcDistributionControlElement IfcEnergyConversionDevice IfcFlowController
IfcOpeningElement IfcAnnotation IfcGrid""".split()}


# ───────────── Revit (OLE compound file → BasicFileInfo, no Revit needed) ─────────────
def read_rvt(p: Path, max_chars: int, **_: Any) -> dict:
    try:
        import olefile
    except ImportError as e:
        raise BimHubError("olefile not installed: pip install -e .") from e
    if not olefile.isOleFile(str(p)):
        raise BimHubError("Not a Revit OLE file (corrupt or cloud-only placeholder not downloaded)")
    with olefile.OleFileIO(str(p)) as ole:
        if not ole.exists("BasicFileInfo"):
            return {"note": "BasicFileInfo stream missing", "open_with": "revit"}
        raw = ole.openstream("BasicFileInfo").read()
        streams = ["/".join(s) for s in ole.listdir()][:60]
    text = raw.decode("utf-16-le", errors="ignore")
    if "Format" not in text and "Build" not in text:
        text = raw.decode("latin-1", errors="ignore")
    clean = re.sub(r"[^\x20-\x7E\u0600-\u06FF\n\r:\\/._-]", " ", text)

    def grab(label: str) -> str | None:
        m = re.search(label + r"\s*:\s*([^\r\n]+)", clean)
        return m.group(1).strip() if m else None
    fmt = grab("Format")
    if not fmt:
        m = re.search(r"Autodesk Revit (\d{4})", clean)
        fmt = m.group(1) if m else None
    return {
        "revit_version": fmt, "build": grab("Build"),
        "worksharing": grab("Worksharing"), "central_model_path": grab("Central Model Path"),
        "username": grab("Username"), "last_save_path": grab("Last Save Path"),
        "is_family": p.suffix.lower() == ".rfa", "ole_streams": streams,
        "open_with": "revit MCP (open the file in Revit; the bridge reads the ACTIVE document)",
    }


# ───────────── Navisworks / DWG: metadata only ─────────────
def read_navis(p: Path, max_chars: int, **_: Any) -> dict:
    with open(p, "rb") as f:
        head = f.read(256)
    return {"format": kind_of(p), "header_hex": head[:16].hex(),
            "open_with": "navisworks MCP (open/append in Navisworks; the add-in reads the open document)",
            "note": "NWC/NWD/NWF are proprietary binaries — content is read through Navisworks, not from the file."}


DWG_VERSIONS = {"AC1015": "2000", "AC1018": "2004", "AC1021": "2007", "AC1024": "2010", "AC1027": "2013",
                "AC1032": "2018"}


def read_dwg(p: Path, max_chars: int, **_: Any) -> dict:
    with open(p, "rb") as f:
        code = f.read(6).decode("ascii", "replace")
    return {"dwg_version_code": code, "autocad_format": DWG_VERSIONS.get(code, "unknown"),
            "open_with": "civil3d MCP / AutoCAD"}


def _e57_xml(f, xml_offset: int, xml_len: int, page: int) -> str:
    """E57 binary sections carry a 4-byte CRC at the end of every page — strip them while reading."""
    out = bytearray()
    pos = xml_offset
    data_per_page = page - 4
    while len(out) < xml_len:
        f.seek(pos)
        in_page = pos % page
        take = min(data_per_page - in_page, xml_len - len(out))
        chunk = f.read(take)
        if not chunk:
            break
        out += chunk
        pos += take
        if pos % page == data_per_page:
            pos += 4
    return out.decode("utf-8", "replace")


def read_point_cloud(p: Path, max_chars: int, **_: Any) -> dict:
    import struct
    ext = p.suffix.lower()
    info: dict[str, Any] = {"format": ext[1:].upper(), "size_mb": round(p.stat().st_size / 1048576, 1),
                            "open_with": "ReCap → RCP, linked in Revit (revit MCP) / Navisworks"}
    try:
        with open(p, "rb") as f:
            head = f.read(375)
            if ext == ".e57" and head[:8] == b"ASTM-E57":
                major, minor, phys, xml_off, xml_len, page = struct.unpack("<IIQQQQ", head[8:48])
                xml = _e57_xml(f, xml_off, xml_len, page or 1024)
                counts = [int(x) for x in re.findall(r'recordCount="(\d+)"', xml)]
                info.update(version=f"{major}.{minor}", scans=len(counts), points=sum(counts),
                            guid=(re.search(r'<guid[^>]*>([^<]+)<', xml) or [None, None])[1],
                            coordinate_metadata=(re.search(r'<coordinateMetadata[^>]*>([^<]*)<', xml) or [None, ""])[1][:500] or None,
                            scan_names=re.findall(r'<name type="String"><!\[CDATA\[([^\]]*)\]\]>', xml)[:50])
            elif ext in (".las", ".laz") and head[:4] == b"LASF":
                vmaj, vmin = head[24], head[25]
                fmt, rec_len, legacy = head[104], struct.unpack("<H", head[105:107])[0], struct.unpack("<I", head[107:111])[0]
                sx, sy, sz, ox, oy, oz, maxx, minx, maxy, miny, maxz, minz = struct.unpack("<12d", head[131:227])
                pts = legacy
                if (vmaj, vmin) >= (1, 4) and len(head) >= 255:
                    pts = struct.unpack("<Q", head[247:255])[0] or legacy
                info.update(version=f"{vmaj}.{vmin}", point_format=fmt, record_length=rec_len, points=pts,
                            compressed=ext == ".laz",
                            bounds={"min": [minx, miny, minz], "max": [maxx, maxy, maxz]},
                            extent_m=[round(maxx - minx, 3), round(maxy - miny, 3), round(maxz - minz, 3)])
            elif ext == ".pts":
                first = head.split(b"\n", 1)[0].strip()
                if first.isdigit():
                    info["points_first_scan"] = int(first)
            elif ext == ".rcp":
                txt = head.decode("utf-8", "replace")
                info["note"] = "ReCap project — references .rcs scans in the same folder tree"
                info["rcs_files_nearby"] = len(list(p.parent.rglob("*.rcs")))
                if txt.lstrip().startswith("<"):
                    info["header"] = txt[:200]
    except (OSError, struct.error) as e:
        info["warning"] = f"header not readable: {e}"
    return info


READERS: dict[str, Callable[..., dict]] = {
    "point-cloud": read_point_cloud,
    "text": read_text, "csv": read_csv, "word": read_docx, "excel": read_xlsx, "pdf": read_pdf,
    "powerpoint": read_pptx, "ifc": read_ifc, "revit-model": read_rvt, "revit-family": read_rvt,
    "revit-template": read_rvt, "navisworks-cache": read_navis, "navisworks-document": read_navis,
    "navisworks-fileset": read_navis, "dwg": read_dwg, "dxf": read_text,
}


def read_file(p: Path, max_chars: int = 30000, **opts: Any) -> dict:
    k = kind_of(p)
    fn = READERS.get(k)
    if not fn:
        raise BimHubError(f"No reader for {p.suffix or 'this file'} ({k}). Legacy .doc/.xls: save as .docx/.xlsx.")
    out = fn(p, max_chars, **opts)
    return {"path": str(p), "kind": k, "size_kb": p.stat().st_size // 1024, **out}
